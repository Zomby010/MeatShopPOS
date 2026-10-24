"""Product import from CSV or Excel, using the same columns as the Bizkit
template so existing product lists can be reused."""

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction

from inventory.models import StockMovement
from inventory.services import set_opening_stock

from .models import Category, Product, UnitOfMeasure
from .services import save_product

COLUMNS = [
    "Product Name",
    "Category",
    "Buying Price",
    "Retail Price",
    "Wholesale Price",
    "Current Stock Level",
    "Low Stock Level At",
    "Units Of Measure",
    "Barcode",
    "SKU",
]

UNIT_ALIASES = {
    "kg": "kg", "kgs": "kg", "kilo": "kg", "kilos": "kg", "kilogram": "kg", "kilograms": "kg",
    "piece": "piece", "pieces": "piece", "pcs": "piece", "pc": "piece", "each": "piece",
    "packet": "packet", "packets": "packet", "pkt": "packet", "pack": "packet", "packs": "packet",
    "tray": "tray", "trays": "tray",
}


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    opening_stock: int = 0
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


def _key(header):
    return " ".join(str(header or "").split()).lower()


def read_rows(uploaded_name, data):
    """Yield dicts keyed by normalised column name from CSV or XLSX bytes."""
    if uploaded_name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = [_key(h) for h in next(rows, [])]
        for values in rows:
            if any(v not in (None, "") for v in values):
                yield dict(zip(headers, values))
    else:
        text = data.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        headers = [_key(h) for h in next(reader, [])]
        for values in reader:
            if any(v.strip() for v in values):
                yield dict(zip(headers, values))


def _decimal(value, label, default=Decimal(0)):
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        raise ValueError(f"{label} '{value}' is not a number")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores barcodes as numbers
    return str(value).strip()


def import_products(uploaded_name, data, *, branch, user=None):
    """Create or update products by name. A stock level is recorded as opening
    stock at `branch` only for products with no stock history there yet."""
    result = ImportResult()
    units = {u.code: u for u in UnitOfMeasure.objects.all()}
    for line_no, row in enumerate(read_rows(uploaded_name, data), start=2):
        name = _text(row.get("product name"))
        if not name:
            result.errors.append(f"Row {line_no}: missing Product Name")
            continue
        try:
            with transaction.atomic():
                unit_code = UNIT_ALIASES.get(_text(row.get("units of measure")).lower(), None)
                if not unit_code or unit_code not in units:
                    raise ValueError(f"unknown unit '{row.get('units of measure')}'")
                category_name = _text(row.get("category")) or "Uncategorised"
                category, _ = Category.objects.get_or_create(
                    name=category_name, defaults={"created_by": user}
                )
                product = Product.objects.filter(name__iexact=name).first()
                is_new = product is None
                product = product or Product(name=name)
                product.category = category
                product.unit = units[unit_code]
                product.is_weighed = units[unit_code].allows_decimal
                product.buying_price = _decimal(row.get("buying price"), "Buying Price")
                product.retail_price = _decimal(row.get("retail price"), "Retail Price")
                product.wholesale_price = _decimal(row.get("wholesale price"), "Wholesale Price")
                product.low_stock_level = _decimal(row.get("low stock level at"), "Low Stock Level At")
                product.barcode = _text(row.get("barcode")) or None
                product.sku = _text(row.get("sku")) or None
                save_product(product, user)
                result.created += is_new
                result.updated += not is_new

                stock = _decimal(row.get("current stock level"), "Current Stock Level")
                if stock:
                    if StockMovement.objects.filter(product=product, branch=branch).exists():
                        result.warnings.append(
                            f"Row {line_no}: {name} already has stock at {branch}; "
                            "stock level ignored (use a recount)."
                        )
                    else:
                        set_opening_stock(
                            product=product, branch=branch, quantity=stock, user=user,
                            reason="Opening stock (import)",
                        )
                        result.opening_stock += 1
        except Exception as exc:  # report the row and carry on with the rest
            result.errors.append(f"Row {line_no} ({name}): {exc}")
    return result


def template_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(COLUMNS)
    ws.append(["Beef Mince", "Beef", 620, 750, 700, 40, 5, "kg", "", "BEEF-MINCE"])
    ws.append(["Eggs (Tray of 30)", "Eggs", 380, 450, 420, 20, 5, "tray", "", "EGGS-TRAY"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
